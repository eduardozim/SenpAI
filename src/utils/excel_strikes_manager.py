"""
Módulo de Exportação e Importação de Golpes em Planilhas Excel (.xlsx) para Treinamento e Calibração.
Fornece utilitários para gerar planilhas profissionais com os eventos detectados,
importar dados revisados por árbitros (Shinpans) e executar o ciclo de retreinamento de golpes.
"""

import io
import os
import re
import datetime
from typing import Dict, Any, List, Tuple, Optional, Union
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.engine.feedback_manager import FeedbackManager, DAN_NAMES, SHINPAN_REV_KEY, SHINPAN_NAME, is_shinpan_reviewer, DuplicateShinpanReviewError
from src.engine.calibrator import CalibrationEngine
from src.utils.logger_manager import log_event


VALID_STRIKE_TYPES = ["MEN", "KOTE", "DO", "TSUKI"]

COLUMN_MAP = {
    "id_evento": "event_id",
    "id": "event_id",
    "event_id": "event_id",
    "evento": "event_id",
    "timestamp": "timestamp",
    "tempo": "timestamp",
    "tempo_formatado": "timestamp",
    "time": "timestamp",
    "tempo_segundos": "time_sec",
    "tempo_s": "time_sec",
    "time_sec": "time_sec",
    "segundos": "time_sec",
    "golpe": "strike_type",
    "tipo": "strike_type",
    "tecnica": "strike_type",
    "strike_type": "strike_type",
    "tipo_de_golpe": "strike_type",
    "lutador_id": "attacker_id",
    "atacante_id": "attacker_id",
    "attacker_id": "attacker_id",
    "lutador": "attacker_id",
    "atacante": "attacker_id",
    "lutador_nome": "attacker_name",
    "atacante_nome": "attacker_name",
    "attacker_name": "attacker_name",
    "ippon_valido": "is_valid_ippon",
    "valido": "is_valid_ippon",
    "is_valid": "is_valid_ippon",
    "ippon": "is_valid_ippon",
    "e_ippon": "is_valid_ippon",
    "is_valid_ippon": "is_valid_ippon",
    "categoria_decisao": "decision_category",
    "categoria": "decision_category",
    "decision_category": "decision_category",
    "rotulo_avaliacao": "label",
    "rotulo": "label",
    "label": "label",
    "pontuacao_total": "total_score",
    "pontuacao": "total_score",
    "score": "total_score",
    "total_score": "total_score",
    "impacto_alvo": "target_impact",
    "impacto": "target_impact",
    "target_impact": "target_impact",
    "fumikomi_sync": "fumikomi_sync",
    "fumikomi": "fumikomi_sync",
    "postura": "posture",
    "posture": "posture",
    "zanshin": "zanshin",
    "origem": "source",
    "source": "source",
    "observacoes": "notes",
    "observacao": "notes",
    "notas": "notes",
    "notes": "notes",
    "comentario": "notes",
    "link_streaming": "streaming_url",
    "link": "streaming_url",
    "url": "streaming_url",
    "streaming_url": "streaming_url",
    "streaming": "streaming_url",
    "url_video": "streaming_url",
    "link_video": "streaming_url",
    "dan_revisor": "reviewer_dan",
    "dan": "reviewer_dan",
    "reviewer_dan": "reviewer_dan",
    "revisor_dan": "reviewer_dan",
    "graduacao_dan": "reviewer_dan"
}


def _clean_column_name(col_name: str) -> str:
    """Normaliza nome de coluna para minúsculas, sem acentos e sem caracteres especiais."""
    s = col_name.lower().strip()
    s = re.sub(r"[áàãâä]", "a", s)
    s = re.sub(r"[éèêë]", "e", s)
    s = re.sub(r"[íìîï]", "i", s)
    s = re.sub(r"[óòõôö]", "o", s)
    s = re.sub(r"[úùûü]", "u", s)
    s = re.sub(r"[ç]", "c", s)
    s = re.sub(r"[^\w\s]", "_", s)
    s = re.sub(r"\s+", "_", s)
    return s.strip("_")


def _format_seconds_to_ts(sec: float) -> str:
    """Converte segundos float para string no formato MM:SS.mmm."""
    sec = max(0.0, sec)
    mins = int(sec // 60)
    remainder = sec % 60
    return f"{mins:02d}:{remainder:06.3f}"


def _parse_ts_to_seconds(ts_str: str) -> float:
    """Converte string timestamp (MM:SS.mmm ou SS.s) para segundos float."""
    if not ts_str:
        return 0.0
    s_val = ts_str.strip()
    if ":" in s_val:
        parts = s_val.split(":")
        try:
            mins = float(parts[0])
            secs = float(parts[1])
            return mins * 60.0 + secs
        except Exception:
            return 0.0
    try:
        return float(s_val)
    except Exception:
        return 0.0


def export_strikes_to_excel(
    strikes_data: List[Dict[str, Any]],
    video_name: str = "match_recorded.mp4",
    metadata: Optional[Dict[str, Any]] = None,
    streaming_url: Optional[str] = None,
    reviewer_dan: Optional[Union[int, str]] = None
) -> bytes:
    """
    Exporta a lista de golpes (detectados pela IA e incluídos pelo revisor) para uma planilha Excel (.xlsx)
    com estilo visual profissional, aba de metadados do vídeo e aba de instruções para edição.
    Inclui obrigatoriamente o Link de Streaming e o Dan do Revisor.
    Retorna os bytes do arquivo .xlsx gerado.
    """
    st_url = str(streaming_url or (metadata.get("streaming_url") if metadata else "") or "").strip()
    dan_val = reviewer_dan if reviewer_dan is not None else (metadata.get("reviewer_dan") if metadata else None)
    if dan_val is None:
        dan_val = 3

    if is_shinpan_reviewer(dan_val):
        dan_label = SHINPAN_NAME
        dan_numeric_str = SHINPAN_REV_KEY
    else:
        try:
            dan_int = int(dan_val)
        except (ValueError, TypeError):
            dan_int = 3
        dan_int = max(1, min(8, dan_int))
        dan_label = DAN_NAMES.get(dan_int, f"{dan_int}º Dan")
        dan_numeric_str = str(dan_int)

    rows = []
    for idx, s in enumerate(strikes_data):
        rev = s.get("review", {})
        raw_ev = s.get("raw_event") or {}
        eval_info = raw_ev.get("evaluation", {})

        is_ippon = rev.get("is_valid_ippon")
        if is_ippon is None:
            is_ippon = s.get("orig_is_valid", eval_info.get("is_valid", False))

        strike_type = (rev.get("strike_type") or s.get("strike_type") or "MEN").strip().upper()
        if strike_type not in VALID_STRIKE_TYPES:
            # Fallback para tipos aceitos
            for st_valid in VALID_STRIKE_TYPES:
                if st_valid in strike_type:
                    strike_type = st_valid
                    break

        att_id = rev.get("attacker_id") or s.get("attacker_id", "KENSHI_AKA")
        att_name = rev.get("attacker_name") or s.get("attacker_label", "Kenshi Aka (Vermelho)")
        ts_val = rev.get("timestamp") or s.get("timestamp", "00:00.000")
        t_sec = s.get("time_sec", _parse_ts_to_seconds(ts_val))

        tot_score = rev.get("total_score")
        if tot_score is None or tot_score == 0.0:
            tot_score = eval_info.get("total_score", 100.0 if is_ippon else 0.0)

        sub_scores = rev.get("sub_scores") or eval_info.get("sub_scores") or {}

        source = s.get("source", "AI_DETECTED")
        source_label = "IA_DETECTADO" if source == "AI_DETECTED" else "INCLUIDO_MANUAL"

        notes = rev.get("notes", "")

        rows.append({
            "ID_Evento": s.get("event_id", f"event_{idx+1}"),
            "Timestamp": ts_val,
            "Tempo_Segundos": round(float(t_sec), 3),
            "Golpe": strike_type,
            "Lutador_ID": att_id,
            "Lutador_Nome": att_name,
            "Ippon_Valido": "SIM" if is_ippon else "NÃO",
            "Categoria_Decisao": "VALID_IPPON" if is_ippon else "INVALID_HIT",
            "Rotulo_Avaliacao": rev.get("label", "TP" if is_ippon else "FP"),
            "Pontuacao_Total": round(float(tot_score), 1),
            "Impacto_Alvo": round(float(sub_scores.get("target_impact", 0.0)), 1),
            "Fumikomi_Sync": round(float(sub_scores.get("fumikomi_sync", 0.0)), 1),
            "Postura": round(float(sub_scores.get("posture", 0.0)), 1),
            "Zanshin": round(float(sub_scores.get("zanshin", 0.0)), 1),
            "Origem": source_label,
            "Dan_Revisor": dan_label,
            "Link_Streaming": st_url,
            "Observacoes": notes
        })

    df = pd.DataFrame(rows)

    # Se a lista estiver vazia, cria DataFrame com colunas corretas
    if df.empty:
        df = pd.DataFrame(columns=[
            "ID_Evento", "Timestamp", "Tempo_Segundos", "Golpe", "Lutador_ID", "Lutador_Nome",
            "Ippon_Valido", "Categoria_Decisao", "Rotulo_Avaliacao", "Pontuacao_Total",
            "Impacto_Alvo", "Fumikomi_Sync", "Postura", "Zanshin", "Origem",
            "Dan_Revisor", "Link_Streaming", "Observacoes"
        ])

    # Metadados do Vídeo e da Sessão de Revisão
    meta_rows = [
        {"Metadado": "Link de Streaming", "Valor": st_url if st_url else "N/A"},
        {"Metadado": "Graduação Dan do Revisor", "Valor": dan_label},
        {"Metadado": "Dan do Revisor (Numérico)", "Valor": dan_numeric_str},
        {"Metadado": "Nome / Arquivo do Vídeo", "Valor": video_name},
        {"Metadado": "Data e Hora da Exportação", "Valor": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {"Metadado": "Total de Golpes Registrados", "Valor": str(len(rows))},
        {"Metadado": "Golpes Válidos (Ippon)", "Valor": str(sum(1 for r in rows if r["Ippon_Valido"] == "SIM"))},
        {"Metadado": "Golpes Inválidos", "Valor": str(sum(1 for r in rows if r["Ippon_Valido"] == "NÃO"))},
        {"Metadado": "Sistema Emissor", "Valor": "SenpAI Kendo Vision System v2.0"}
    ]
    df_meta = pd.DataFrame(meta_rows)

    # Instruções e Dicionário na terceira aba
    instructions_data = [
        {"Coluna": "ID_Evento", "Obrigatorio": "Não", "Valores Permitidos": "Texto livre ou em branco", "Descricao": "Identificador único do evento. Deixe em branco se estiver adicionando uma nova linha de golpe não detectado."},
        {"Coluna": "Timestamp", "Obrigatorio": "Sim", "Valores Permitidos": "MM:SS.mmm ou segundos (ex: 00:04.120 ou 4.12)", "Descricao": "Momento exato do golpe no vídeo."},
        {"Coluna": "Golpe", "Obrigatorio": "Sim", "Valores Permitidos": "MEN, KOTE, DO, TSUKI", "Descricao": "Técnica de ataque desferida pelo kendoca."},
        {"Coluna": "Lutador_ID", "Obrigatorio": "Sim", "Valores Permitidos": "KENSHI_AKA ou KENSHI_SHIRO", "Descricao": "Identificação do kendoca que atacou (Vermelho ou Branco)."},
        {"Coluna": "Ippon_Valido", "Obrigatorio": "Sim", "Valores Permitidos": "SIM ou NÃO", "Descricao": "Indica se o golpe atingiu os requisitos de Yuko-Datotsu (Ki-Ken-Tai-Ichi e Zanshin) para ser ponto válido (Ippon)."},
        {"Coluna": "Categoria_Decisao", "Obrigatorio": "Não", "Valores Permitidos": "VALID_IPPON ou INVALID_HIT", "Descricao": "Categoria técnica formal do golpe."},
        {"Coluna": "Rotulo_Avaliacao", "Obrigatorio": "Não", "Valores Permitidos": "TP (Golpe Válido), FP (Inválido/Falso Positivo), FN (Golpe Perdido)", "Descricao": "Rótulo de validação supervisionada."},
        {"Coluna": "Pontuacao_Total", "Obrigatorio": "Não", "Valores Permitidos": "0.0 a 100.0", "Descricao": "Pontuação ponderada global do golpe."},
        {"Coluna": "Dan_Revisor", "Obrigatorio": "Não", "Valores Permitidos": "1º Dan a 8º Dan", "Descricao": "Graduação Dan do árbitro (Shinpan) ou revisor técnico responsável pela avaliação."},
        {"Coluna": "Link_Streaming", "Obrigatorio": "Sim", "Valores Permitidos": "URL válida de streaming (YouTube / Web Stream)", "Descricao": "Link de streaming da partida de Kendo de onde os eventos foram extraídos."},
        {"Coluna": "Observacoes", "Obrigatorio": "Não", "Valores Permitidos": "Texto livre", "Descricao": "Anotações técnicas do árbitro/avaliador para justificar a pontuação ou correção."}
    ]
    df_instructions = pd.DataFrame(instructions_data)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Golpes Detectados", index=False)
        df_meta.to_excel(writer, sheet_name="Metadados & Vídeo", index=False)
        df_instructions.to_excel(writer, sheet_name="Instruções & Dicionário", index=False)

        workbook = writer.book
        ws_strikes = writer.sheets["Golpes Detectados"]
        ws_meta = writer.sheets["Metadados & Vídeo"]
        ws_inst = writer.sheets["Instruções & Dicionário"]

        # Estilização da Aba 1 (Golpes Detectados)
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        ws_strikes.row_dimensions[1].height = 28
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws_strikes.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Cores para Ippon Válido
        valid_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Verde suave
        invalid_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Vermelho suave

        for row_idx in range(2, len(df) + 2):
            ws_strikes.row_dimensions[row_idx].height = 20
            is_valid_val = str(ws_strikes.cell(row=row_idx, column=7).value or "").upper()
            row_fill = valid_fill if is_valid_val in ["SIM", "TRUE", "VALID_IPPON"] else (invalid_fill if is_valid_val in ["NÃO", "NAO", "FALSE", "INVALID_HIT"] else None)

            for col_idx in range(1, len(df.columns) + 1):
                c = ws_strikes.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                c.font = Font(name="Segoe UI", size=10)
                if col_idx in [1, 2, 4, 5, 7, 8, 9, 15, 16]:
                    c.alignment = center_align
                elif col_idx in [3, 10, 11, 12, 13, 14]:
                    c.alignment = right_align
                else:
                    c.alignment = left_align

                # Destaque condicional sutil na coluna Ippon_Valido
                if col_idx == 7 and row_fill:
                    c.fill = row_fill
                    c.font = Font(name="Segoe UI", size=10, bold=True)

        # Ajuste de largura de colunas (Aba 1)
        for col in ws_strikes.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_strikes.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Estilização da Aba 2 (Metadados & Vídeo)
        ws_meta.row_dimensions[1].height = 26
        meta_header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
        for col_idx in range(1, len(df_meta.columns) + 1):
            cell = ws_meta.cell(row=1, column=col_idx)
            cell.fill = meta_header_fill
            cell.font = header_font
            cell.alignment = center_align

        for row_idx in range(2, len(df_meta) + 2):
            ws_meta.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(df_meta.columns) + 1):
                c = ws_meta.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                c.font = Font(name="Segoe UI", size=10, bold=(col_idx == 1))
                c.alignment = left_align

        for col in ws_meta.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_meta.column_dimensions[col_letter].width = max(max_len + 6, 22)

        # Estilização da Aba 3 (Instruções)
        ws_inst.row_dimensions[1].height = 26
        inst_header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        for col_idx in range(1, len(df_instructions.columns) + 1):
            cell = ws_inst.cell(row=1, column=col_idx)
            cell.fill = inst_header_fill
            cell.font = header_font
            cell.alignment = center_align

        for row_idx in range(2, len(df_instructions) + 2):
            ws_inst.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(df_instructions.columns) + 1):
                c = ws_inst.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                c.font = Font(name="Segoe UI", size=10)
                c.alignment = center_align if col_idx in [1, 2] else left_align

        for col in ws_inst.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_inst.column_dimensions[col_letter].width = max(max_len + 5, 15)

    return output.getvalue()


def import_strikes_from_excel(
    file_content_or_path: Union[str, bytes, io.BytesIO],
    existing_events: Optional[List[Dict[str, Any]]] = None
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Importa e valida os golpes de uma planilha Excel (.xlsx ou .xls).
    Permite atualizar golpes existentes e incluir novos golpes adicionados manualmente na planilha.
    Retorna:
      (lista_de_golpes_normalizados, dicionario_resumo_estatistico)
    """
    meta_streaming_url = ""
    meta_reviewer_dan = None

    # Tenta extrair metadados da aba dedicada 'Metadados & Vídeo' caso exista
    try:
        if isinstance(file_content_or_path, io.BytesIO):
            file_content_or_path.seek(0)
        df_meta = pd.read_excel(file_content_or_path, sheet_name="Metadados & Vídeo")
        if not df_meta.empty and len(df_meta.columns) >= 2:
            col_k = df_meta.columns[0]
            col_v = df_meta.columns[1]
            for _, m_row in df_meta.iterrows():
                k_txt = _clean_column_name(str(m_row.get(col_k, "")))
                v_str = str(m_row.get(col_v, "")).strip()
                if any(term in k_txt for term in ["streaming", "link", "url"]):
                    if v_str and v_str not in ["nan", "None", "N/A", "Nenhum link informado"]:
                        meta_streaming_url = v_str
                elif "dan" in k_txt or "revisor" in k_txt or "arbitr" in k_txt:
                    if is_shinpan_reviewer(v_str):
                        meta_reviewer_dan = SHINPAN_REV_KEY
                    else:
                        num_m = re.search(r"\d+", v_str)
                        if num_m:
                            meta_reviewer_dan = int(num_m.group(0))
    except Exception:
        pass

    try:
        # Tenta ler a aba principal "Golpes Detectados", se não existir lê a primeira aba
        if isinstance(file_content_or_path, io.BytesIO):
            file_content_or_path.seek(0)
        try:
            df = pd.read_excel(file_content_or_path, sheet_name="Golpes Detectados")
        except Exception:
            if isinstance(file_content_or_path, io.BytesIO):
                file_content_or_path.seek(0)
            df = pd.read_excel(file_content_or_path, sheet_name=0)
    except Exception as e:
        raise ValueError(f"Não foi possível abrir o arquivo Excel: {str(e)}")

    if df.empty:
        return [], {
            "status": "empty",
            "total_rows": 0,
            "valid_ippons": 0,
            "invalid_hits": 0,
            "new_strikes_count": 0,
            "warnings": ["A planilha importada está vazia."],
            "streaming_url": meta_streaming_url,
            "reviewer_dan": meta_reviewer_dan,
            "reviewer_dan_name": DAN_NAMES.get(meta_reviewer_dan, f"{meta_reviewer_dan}º Dan") if meta_reviewer_dan else None
        }

    # Mapeamento e normalização dos nomes das colunas
    normalized_cols = {}
    for col in df.columns:
        cleaned = _clean_column_name(col)
        if cleaned in COLUMN_MAP:
            normalized_cols[col] = COLUMN_MAP[cleaned]
        else:
            # Busca por substring
            matched = False
            for k, v in COLUMN_MAP.items():
                if k in cleaned or cleaned in k:
                    normalized_cols[col] = v
                    matched = True
                    break
            if not matched:
                normalized_cols[col] = cleaned

    df = df.rename(columns=normalized_cols)

    # Se streaming_url ou reviewer_dan não estavam na aba de metadados, procura nas colunas de df
    if not meta_streaming_url:
        for c in ["streaming_url", "link_streaming", "link", "url"]:
            if c in df.columns:
                non_empty = df[c].dropna()
                for val in non_empty:
                    v_str = str(val).strip()
                    if v_str and v_str not in ["nan", "None", "N/A"]:
                        meta_streaming_url = v_str
                        break
                if meta_streaming_url:
                    break

    if meta_reviewer_dan is None:
        for c in ["reviewer_dan", "dan_revisor", "dan"]:
            if c in df.columns:
                non_empty = df[c].dropna()
                for val in non_empty:
                    v_str = str(val).strip()
                    if is_shinpan_reviewer(v_str):
                        meta_reviewer_dan = SHINPAN_REV_KEY
                        break
                    num_m = re.search(r"\d+", v_str)
                    if num_m:
                        meta_reviewer_dan = int(num_m.group(0))
                        break
                if meta_reviewer_dan is not None:
                    break

    parsed_strikes: List[Dict[str, Any]] = []
    warnings: List[str] = []
    new_strikes_count = 0
    valid_ippons_count = 0
    invalid_hits_count = 0

    existing_ids = set()
    if existing_events:
        for ev in existing_events:
            ev_id = ev.get("event_id") or (ev.get("review", {}).get("event_id") if isinstance(ev.get("review"), dict) else None)
            if ev_id:
                existing_ids.add(str(ev_id).strip())

    for row_pos, (idx, row) in enumerate(df.iterrows()):
        # Ignora linhas totalmente vazias
        if row.isna().all():
            continue

        raw_id = str(row.get("event_id", "")).strip()
        if raw_id in ["nan", "None", "", "null"]:
            raw_id = ""

        # 1. Timestamp
        raw_ts = row.get("timestamp")
        raw_sec = row.get("time_sec")
        if pd.isna(raw_ts) or str(raw_ts).strip() in ["nan", "None", ""]:
            if not pd.isna(raw_sec):
                ts_str = _format_seconds_to_ts(float(raw_sec))
                t_sec = float(raw_sec)
            else:
                warnings.append(f"Linha {row_pos+2}: Timestamp ausente. Ignorando linha.")
                continue
        else:
            ts_val = str(raw_ts).strip()
            if ":" in ts_val:
                ts_str = ts_val
                t_sec = _parse_ts_to_seconds(ts_val)
            else:
                try:
                    s_float = float(ts_val)
                    ts_str = _format_seconds_to_ts(s_float)
                    t_sec = s_float
                except ValueError:
                    ts_str = "00:00.000"
                    t_sec = 0.0
                    warnings.append(f"Linha {row_pos+2}: Formato de timestamp '{ts_val}' inválido. Ajustado para 00:00.000.")

        # 2. Tipo de Golpe (Técnica)
        raw_strike = str(row.get("strike_type", "MEN")).strip().upper()
        if raw_strike in ["nan", "None", ""]:
            raw_strike = "MEN"

        strike_type = "MEN"
        for st_candidate in VALID_STRIKE_TYPES:
            if st_candidate in raw_strike:
                strike_type = st_candidate
                break

        # 3. Validação de Ippon
        raw_val = str(row.get("is_valid_ippon", "")).strip().upper()
        raw_cat = str(row.get("decision_category", "")).strip().upper()

        if raw_val in ["SIM", "S", "TRUE", "1", "V", "VERDADEIRO", "YES", "Y", "VALID_IPPON", "IPPON"]:
            is_valid_ippon = True
        elif raw_val in ["NÃO", "NAO", "N", "FALSE", "0", "F", "FALSO", "NO", "INVALID_HIT"]:
            is_valid_ippon = False
        else:
            is_valid_ippon = (raw_cat == "VALID_IPPON")

        decision_cat = "VALID_IPPON" if is_valid_ippon else "INVALID_HIT"
        label_val = str(row.get("label", "")).strip().upper()
        if is_valid_ippon:
            if label_val not in ["TP", "CONFIRMED", "INCLUDED", "EDITED"]:
                label_val = "TP"
        else:
            if label_val in ["TP", "CONFIRMED"]:
                label_val = "FP"
            elif label_val not in ["FP", "FN", "INCLUDED", "EDITED"]:
                label_val = "FP"

        if is_valid_ippon:
            valid_ippons_count += 1
        else:
            invalid_hits_count += 1

        # 4. Atacante (Aka ou Shiro)
        raw_att_id = str(row.get("attacker_id", "")).strip().upper()
        raw_att_name = str(row.get("attacker_name", "")).strip()

        if any(k in raw_att_id for k in ["SHIRO", "BRANCO", "WHITE", "2"]) or any(k in raw_att_name.upper() for k in ["SHIRO", "BRANCO", "WHITE"]):
            att_id = "KENSHI_SHIRO"
            att_name = "Kenshi Shiro (Branco)"
        else:
            att_id = "KENSHI_AKA"
            att_name = "Kenshi Aka (Vermelho)"

        # 5. Pontuações e Sub-scores
        def _safe_float(val: Any, default: float = 0.0) -> float:
            try:
                if pd.isna(val):
                    return default
                return float(val)
            except Exception:
                return default

        tot_score = _safe_float(row.get("total_score"), 100.0 if is_valid_ippon else 0.0)
        sub_scores = {
            "target_impact": _safe_float(row.get("target_impact"), 85.0 if is_valid_ippon else 40.0),
            "fumikomi_sync": _safe_float(row.get("fumikomi_sync"), 75.0 if is_valid_ippon else 40.0),
            "posture": _safe_float(row.get("posture"), 80.0 if is_valid_ippon else 45.0),
            "zanshin": _safe_float(row.get("zanshin"), 75.0 if is_valid_ippon else 35.0)
        }

        # 6. Identificador do Evento e Origem
        is_new = False
        if not raw_id:
            raw_id = f"fn_{ts_str.replace(':', '_').replace('.', '_')}_{att_id.lower()}_{row_pos+1}"
            is_new = True
            new_strikes_count += 1
        elif existing_ids and raw_id not in existing_ids:
            is_new = True
            new_strikes_count += 1

        notes = str(row.get("notes", "")).strip()
        if notes in ["nan", "None"]:
            notes = ""

        row_dan_raw = row.get("reviewer_dan")
        row_dan = None
        if row_dan_raw is not None and not pd.isna(row_dan_raw):
            if is_shinpan_reviewer(row_dan_raw):
                row_dan = SHINPAN_REV_KEY
            else:
                num_m = re.search(r"\d+", str(row_dan_raw))
                if num_m:
                    row_dan = int(num_m.group(0))
        if row_dan is None:
            row_dan = meta_reviewer_dan

        row_url_raw = str(row.get("streaming_url", "")).strip()
        if row_url_raw in ["nan", "None", "N/A", ""]:
            row_url = meta_streaming_url
        else:
            row_url = row_url_raw

        parsed_strikes.append({
            "event_id": raw_id,
            "timestamp": ts_str,
            "time_sec": t_sec,
            "strike_type": strike_type,
            "attacker_id": att_id,
            "attacker_name": att_name,
            "is_valid_ippon": is_valid_ippon,
            "decision_category": decision_cat,
            "label": label_val,
            "total_score": tot_score,
            "sub_scores": sub_scores,
            "is_included": is_new,
            "is_edited": True,
            "is_confirmed": True,
            "notes": notes,
            "dan_revisor": row_dan,
            "streaming_url": row_url,
            "source": "INCLUDED" if is_new else "EXCEL_EDITED"
        })

    # Ordena cronologicamente
    parsed_strikes.sort(key=lambda s: s["time_sec"])

    summary = {
        "status": "success",
        "total_rows": len(parsed_strikes),
        "valid_ippons": valid_ippons_count,
        "invalid_hits": invalid_hits_count,
        "new_strikes_count": new_strikes_count,
        "warnings": warnings,
        "streaming_url": meta_streaming_url,
        "reviewer_dan": meta_reviewer_dan,
        "reviewer_dan_name": DAN_NAMES.get(meta_reviewer_dan, f"{meta_reviewer_dan}º Dan") if meta_reviewer_dan else None
    }

    return parsed_strikes, summary


def apply_imported_strikes_to_session_reviews(
    imported_strikes: List[Dict[str, Any]],
    current_session_reviews: Optional[Dict[str, Any]] = None
) -> Dict[str, Any]:
    """
    Atualiza ou preenche o dicionário `st.session_state["session_reviews"]` com os dados importados do Excel,
    garantindo que a linha do tempo, o placar oficial e as telas de visualização reflitam imediatamente as edições.
    """
    session_revs = dict(current_session_reviews or {})

    for item in imported_strikes:
        ev_id = item["event_id"]
        session_revs[ev_id] = {
            "event_id": ev_id,
            "label": item.get("label", "TP" if item.get("is_valid_ippon") else "FP"),
            "category": item.get("decision_category", "VALID_IPPON" if item.get("is_valid_ippon") else "INVALID_HIT"),
            "decision_category": item.get("decision_category", "VALID_IPPON" if item.get("is_valid_ippon") else "INVALID_HIT"),
            "is_valid_ippon": bool(item.get("is_valid_ippon")),
            "strike_type": item.get("strike_type", "MEN"),
            "timestamp": item.get("timestamp", "00:00.000"),
            "attacker_id": item.get("attacker_id", "KENSHI_AKA"),
            "attacker_name": item.get("attacker_name", "Kenshi Aka (Vermelho)"),
            "total_score": float(item.get("total_score", 0.0)),
            "sub_scores": item.get("sub_scores", {}),
            "is_included": bool(item.get("is_included", False)),
            "is_confirmed": True,
            "is_edited": True,
            "notes": str(item.get("notes", ""))
        }

    return session_revs


def execute_training_from_imported_strikes(
    imported_strikes: List[Dict[str, Any]],
    video_name: str,
    profile_key: str,
    reviewer_dan: Any,
    current_profile_config: Dict[str, Any],
    feedback_mgr: FeedbackManager,
    auto_trainer_instance: Optional[Any] = None,
    streaming_url: Optional[str] = None
) -> Dict[str, Any]:
    """
    Executa o treinamento adaptativo dos golpes a partir dos dados importados do Excel:
    1. Grava as anotações supervisionadas com governança por Dan ou Shinpans no FeedbackManager.
    2. Recalibra os limiares de decisão de Ki-Ken-Tai-Ichi e Zanshin e os pesos dos golpes.
    3. Atualiza o perfil ativo no CalibrationEngine.
    4. Opcionalmente alimenta a base de conhecimento do AutoTrainer para o escopo de combates gravados (recorded_shiai).
    """
    if is_shinpan_reviewer(reviewer_dan):
        dan_val = SHINPAN_REV_KEY
        dan_name = SHINPAN_NAME
    else:
        try:
            dan_int = int(reviewer_dan)
            dan_val = max(1, min(8, dan_int))
        except Exception:
            dan_val = 1
        dan_name = DAN_NAMES.get(dan_val, f"{dan_val}º Dan")

    # Salva sessão de revisão via feedback manager
    new_cfg, session_rec = feedback_mgr.save_review_session(
        video_name=video_name,
        profile_key=profile_key,
        reviewer_dan=dan_val,
        review_items=imported_strikes,
        current_profile_config=current_profile_config,
        video_url=streaming_url
    )

    # Persiste o novo perfil de calibração
    calibrator = CalibrationEngine()
    calibrator.update_and_save_profile(profile_key, new_cfg)

    # Alimenta o AutoTrainer se fornecido
    auto_trainer_res = None
    if auto_trainer_instance is not None:
        try:
            if hasattr(auto_trainer_instance, "record_review_feedback"):
                auto_trainer_res = auto_trainer_instance.record_review_feedback(
                    scope="recorded_shiai",
                    items_count=len(imported_strikes),
                    reviewer_dan=dan_val
                )
            elif hasattr(auto_trainer_instance, "load_knowledge_base") and hasattr(auto_trainer_instance, "save_knowledge_base"):
                kb = auto_trainer_instance.load_knowledge_base()
                shiai_params = kb.get("learned_parameters", {}).get("shiai_scoring", {})
                # Reforça parâmetros aprendidos
                if "optimal_weights" in shiai_params and "weights" in new_cfg:
                    shiai_params["optimal_weights"] = new_cfg["weights"]
                kb["training_sessions_completed"] = kb.get("training_sessions_completed", 0) + 1
                auto_trainer_instance.save_knowledge_base(kb)
                auto_trainer_res = {"status": "knowledge_base_updated"}
        except Exception as e:
            log_event("WARNING", f"Erro ao integrar auto_trainer no retreinamento por Excel: {e}")

    log_event("INFO", f"Treinamento por Excel concluído: {len(imported_strikes)} golpes processados sob governança de {dan_name}.")

    return {
        "status": "success",
        "items_count": len(imported_strikes),
        "reviewer_dan": dan_val,
        "reviewer_dan_name": dan_name,
        "is_shinpan_decision": session_rec.get("is_shinpan_decision", False),
        "new_config": new_cfg,
        "session_record": session_rec,
        "optimization_summary": session_rec.get("optimization_summary", {}),
        "auto_trainer_res": auto_trainer_res
    }
